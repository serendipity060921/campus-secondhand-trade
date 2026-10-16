# -*- coding: utf-8 -*-
"""生成论文图表包（第2批-⑦）

把项目里散落的实测数据汇总成一个 Excel 工作簿 `docs/论文图表包.xlsx`：
每个主题一个工作表，数据 + **Excel 原生图表**（折线/柱状/饼图）。
用原生图表而不是贴图片，是为了让论文里可以直接双击改样式、也能随数据更新。

数据来源全部是**实测产物**，不手工填数：
  · 覆盖率曲线   ← docs/测试覆盖率说明.md 的 13 个数据点
  · 缓存性能对照 ← docs/性能测试报告-缓存效果对照.md 的实测表格
  · 击穿/多键实验 ← tools/perf-stampede.json、tools/perf-multi.json
  · 机审统计     ← 数据库 moderation_log（真实落库数据）
  · 测试用例分布 ← docs/测试覆盖率说明.md 的用例表
  · 缺陷统计     ← docs/缺陷跟踪表.md
  · 索引审查     ← docs/数据库性能与索引审查-数据.json

用法：
    python tools/gen-paper-figures.py
产物：
    docs/论文图表包.xlsx
"""

# ★ Windows 控制台默认 GBK，脚本会输出 ✓/✗，必须重配置 stdout
import sys as _sys
try:
    _sys.stdout.reconfigure(encoding='utf-8')
    _sys.stderr.reconfigure(encoding='utf-8')
except Exception:
    pass

import json
import os
import pathlib
import re
import subprocess

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parent.parent
DOCS = ROOT / 'docs'
TOOLS = ROOT / 'tools'
OUT = DOCS / '论文图表包.xlsx'
MYSQL = os.environ.get('MYSQL_CLI', r'D:\major\tool\mysql-8.4.4-winx64\bin\mysql.exe')

HEAD_FILL = PatternFill('solid', fgColor='DDEBF7')
TITLE_FONT = Font(bold=True, size=12)
HEAD_FONT = Font(bold=True)


def sql(query):
    r = subprocess.run([MYSQL, '-h', '127.0.0.1', '-u', 'root', '-p123456', '-N', '-B',
                        '--default-character-set=utf8mb4', 'campus_trade', '-e', query],
                       capture_output=True, text=True, encoding='utf-8', timeout=60)
    return [l.split('\t') for l in (r.stdout or '').strip().splitlines() if l]


def clean(s):
    return re.sub(r'[*`]', '', str(s)).strip()


def sheet_header(ws, title, headers, note=''):
    """写标题行 + 表头，返回表头所在行号"""
    ws['A1'] = title
    ws['A1'].font = TITLE_FONT
    if note:
        ws['A2'] = note
        ws['A2'].font = Font(italic=True, color='666666')
        row = 4
    else:
        row = 3
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal='center')
    return row


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w


# ---------------------------------------------------------------- 各数据源

def coverage_series():
    """覆盖率折线：返回 (日期, 本轮补了什么, 累计覆盖率%)"""
    text = (DOCS / '测试覆盖率说明.md').read_text(encoding='utf-8')
    out = []
    for line in text.splitlines():
        if not line.startswith('| 20'):
            continue
        cols = [c.strip() for c in line.strip('|').split('|')]
        if len(cols) < 3:
            continue
        m = re.search(r'(\d+\.\d+)%', cols[2])
        if not m:
            continue
        out.append((cols[0], clean(cols[1])[:80], float(m.group(1))))
    return out


def perf_rows():
    """缓存性能对照：场景 | 开启QPS | 关闭QPS | 开启延迟ms | 关闭延迟ms | 开启DB QPS | 关闭DB QPS"""
    text = (DOCS / '性能测试报告-缓存效果对照.md').read_text(encoding='utf-8')
    out = []
    for line in text.splitlines():
        if not line.startswith('| ') or set(line) <= set('|- '):
            continue
        cols = [clean(c) for c in line.strip('|').split('|')]
        if len(cols) != 8:
            continue
        nums = []
        for idx in (1, 2, 4, 5, 6, 7):
            m = re.search(r'(\d+(?:\.\d+)?)', cols[idx].replace(',', ''))
            if not m:
                nums = []
                break
            nums.append(float(m.group(1)))
        if len(nums) == 6:
            out.append((cols[0], *nums))
    return out


def moderation_stats():
    by_action = sql("SELECT action, COUNT(*) FROM moderation_log GROUP BY action ORDER BY action;")
    trend = sql("SELECT DATE_FORMAT(create_time, '%Y-%m-%d') d, "
                "SUM(action='PASS'), SUM(action='REVIEW'), SUM(action='REJECT') "
                "FROM moderation_log GROUP BY d ORDER BY d;")
    rules = sql("SELECT IFNULL(risk_rule,'(无)'), COUNT(*) FROM moderation_log "
                "WHERE risk_rule IS NOT NULL AND risk_rule <> '' GROUP BY risk_rule "
                "ORDER BY COUNT(*) DESC LIMIT 8;")
    return by_action, trend, rules


def case_distribution():
    text = (DOCS / '测试覆盖率说明.md').read_text(encoding='utf-8')
    out = []
    for line in text.splitlines():
        m = re.match(r'^\|\s*\**([\u4e00-\u9fa5A-Za-z ]+)\**\s*\|\s*\**(\d+)\**\s*\|', line)
        if m and m.group(1).strip() not in ('合计', '类别'):
            out.append((m.group(1).strip(), int(m.group(2))))
    return out


def defect_stats():
    """按【章节】统计缺陷数，并与文档里声明的数字交叉核对。

    ★ 为什么不能按编号区间统计：第一版用 `BUG-00[1-9]` 之类区间去数，
      结果把"产品缺陷 6 项"数成了 12 项 —— 编号是连续的，但**分类不是按编号切的**
      （BUG-011/012 属于产品缺陷，BUG-013 起又回到工程缺陷）。
      改成按 `## 一、产品缺陷` / `## 二、工程与测试缺陷` 两个章节分别计数后准确。
      同时交叉核对文档统计表里的声明值：不一致就打印警告 —— 本次正是靠这一点
      发现缺陷表自己的统计数字漏更新了（加了 BUG-017 却没改"10 项"）。
    """
    text = (DOCS / '缺陷跟踪表.md').read_text(encoding='utf-8')
    product = engineering = 0
    sev = {}
    for sec in re.split(r'^## ', text, flags=re.M):
        title = (sec.splitlines() or [''])[0]
        n = len(re.findall(r'^\| BUG-', sec, re.M))
        if title.startswith('一、产品缺陷'):
            product = n
        elif title.startswith('二、工程与测试缺陷'):
            engineering = n
        for line in sec.splitlines():
            if line.startswith('| BUG-'):
                cols = [clean(c) for c in line.strip('|').split('|')]
                if len(cols) >= 4:
                    sev[cols[3]] = sev.get(cols[3], 0) + 1

    declared = {}
    for line in text.splitlines():
        m = re.match(r'^\|\s*(产品缺陷|工程与测试缺陷)\s*\|\s*\**(\d+)\s*项', line)
        if m:
            declared[m.group(1)] = int(m.group(2))
    if declared:
        d_prod = declared.get('产品缺陷')
        d_eng = declared.get('工程与测试缺陷')
        if d_prod != product or d_eng != engineering:
            print(f'  ⚠️ 缺陷统计交叉核对不一致：'
                  f'实际（按章节数）产品 {product} / 工程 {engineering}，'
                  f'而文档统计表声明 产品 {d_prod} / 工程 {d_eng} —— 请同步更新 docs/缺陷跟踪表.md 的统计表')
    return product, engineering, sev


def load_json(name):
    p = TOOLS / name
    if not p.exists():
        p = DOCS / name
    return json.loads(p.read_text(encoding='utf-8')) if p.exists() else {}


# ---------------------------------------------------------------- 建工作簿

def build():
    wb = Workbook()

    # ---- 0. 说明页 ----
    ws = wb.active
    ws.title = '说明'
    ws['A1'] = '校园二手交易平台 · 论文图表包'
    ws['A1'].font = Font(bold=True, size=14)
    lines = [
        '',
        '本工作簿由 tools/gen-paper-figures.py 自动生成，数据全部来自实测产物，未手工填数。',
        '图表为 Excel 原生图表（折线/柱状/饼图），可直接复制进 Word，也可双击改样式。',
        '',
        '工作表                内容                                        数据来源',
        '覆盖率提升曲线        13 个数据点（3.38% → 37.25%）               docs/测试覆盖率说明.md',
        '缓存性能对照          开启/关闭缓存各场景 QPS、延迟、DB QPS        docs/性能测试报告-缓存效果对照.md',
        '击穿实验              单键 50 并发与多热点键（锁粒度）实测         tools/perf-stampede.json / perf-multi.json',
        '机审统计              放行/转审/拒绝分布、按天趋势、命中规则       数据库 moderation_log（真实落库）',
        '测试用例分布          五类自动化用例共 512 条                     docs/测试覆盖率说明.md',
        '缺陷统计              产品缺陷与工程/测试缺陷                     docs/缺陷跟踪表.md',
        '索引审查              21 种查询形态的 EXPLAIN 结果               docs/数据库性能与索引审查-数据.json',
    ]
    for i, t in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=t)
    ws.column_dimensions['A'].width = 120

    # ---- 1. 覆盖率提升曲线 ----
    cov = coverage_series()
    ws = wb.create_sheet('覆盖率提升曲线')
    head = sheet_header(ws, '测试覆盖率提升曲线（JaCoCo 指令覆盖率）', ['日期', '本轮补的测试', '累计覆盖率(%)'],
                        '数据点来自 docs/测试覆盖率说明.md；每次补完测试就把棘轮阈值上调一次，只许上升。')
    for i, (d, what, pct) in enumerate(cov, start=head + 1):
        ws.cell(row=i, column=1, value=d)
        ws.cell(row=i, column=2, value=what)
        ws.cell(row=i, column=3, value=pct)
    if cov:
        ch = LineChart()
        ch.title = '测试覆盖率提升曲线'
        ch.y_axis.title = '覆盖率 (%)'
        ch.x_axis.title = '数据点序号'
        ch.height, ch.width = 9, 20
        data = Reference(ws, min_col=3, min_row=head, max_row=head + len(cov))
        ch.add_data(data, titles_from_data=True)
        ws.add_chart(ch, f'E{head}')
    autosize(ws, [12, 60, 14])

    # ---- 2. 缓存性能对照 ----
    pr = perf_rows()
    ws = wb.create_sheet('缓存性能对照')
    head = sheet_header(
        ws, '缓存效果对照（同一压测脚本、同一并发，仅切换缓存开关）',
        ['场景', '开启缓存 QPS', '关闭缓存 QPS', '开启 平均延迟(ms)', '关闭 平均延迟(ms)',
         '开启 DB QPS', '关闭 DB QPS'],
        '数据来自 docs/性能测试报告-缓存效果对照.md 的实测表格；QPS 越高越好，DB QPS 越低越好。')
    for i, row in enumerate(pr, start=head + 1):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    if pr:
        ch = BarChart()
        ch.type = 'col'
        ch.title = '各场景 QPS：开启 vs 关闭缓存'
        ch.y_axis.title = 'QPS'
        ch.height, ch.width = 9, 22
        ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=head, max_row=head + len(pr)),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=head + len(pr)))
        ws.add_chart(ch, f'I{head}')
    autosize(ws, [22, 13, 13, 16, 16, 12, 12])

    # ---- 3. 击穿实验 ----
    st = load_json('perf-stampede.json')
    mu = load_json('perf-multi.json')
    ws = wb.create_sheet('击穿实验')
    head = sheet_header(ws, '缓存击穿实验（热点 key 过期瞬间）', ['实验', '设置', '关键结果', '结论'],
                        '数据来自 tools/perf-stampede.json 与 tools/perf-multi.json。')
    rows = []
    if st:
        rows.append(('单键击穿', f"商品 {st.get('product_id')}，{st.get('concurrency')} 并发",
                     f"重建 {st.get('rebuilds_avg')} 次（期望 {st.get('naive_expectation')} 次），"
                     f"保护倍数 {st.get('protection_ratio')}×",
                     '互斥重建生效，只重建 1 次'))
    if mu:
        rows.append(('多热点键', f"{mu.get('key_count')} 个 key × 每键 {mu.get('per_key')} 并发"
                                 f"（总并发 {mu.get('total_concurrency')}）",
                     f"多键耗时/单键耗时 = {mu.get('avg_elapsed_ratio')}；"
                     f"合计重建 {mu.get('avg_rebuilds_from_log')} 次（期望 {mu.get('expected_rebuilds')}）",
                     '锁按 key 分（可并行重建），非全局锁'))
    for i, r in enumerate(rows, start=head + 1):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v)
    # 单键逐轮明细
    r0 = head + len(rows) + 3
    ws.cell(row=r0, column=1, value='单键逐轮明细').font = HEAD_FONT
    for j, h in enumerate(['轮次', '请求数', 'DB 语句增量', '重建次数', '总耗时(ms)', 'P95(ms)'], start=1):
        c = ws.cell(row=r0 + 1, column=j, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for i, d in enumerate(st.get('details', []), start=r0 + 2):
        vals = [d.get('round'), d.get('requests'), d.get('db_questions_delta'),
                d.get('rebuilds'), d.get('elapsed_ms'), d.get('p95_ms')]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v)
    # 多键逐轮明细
    r1 = r0 + len(st.get('details', [])) + 3
    ws.cell(row=r1, column=1, value='多热点键逐轮明细（含逐 key 精确重建次数）').font = HEAD_FONT
    for j, h in enumerate(['轮次', '单键耗时(ms)', '多键耗时(ms)', 'TN/T1', '日志逐 key 重建次数', '合计'], start=1):
        c = ws.cell(row=r1 + 1, column=j, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for i, d in enumerate(mu.get('details', []), start=r1 + 2):
        mk = d.get('multi_key', {})
        vals = [d.get('round'), d.get('single_key', {}).get('elapsed_ms'), mk.get('elapsed_ms'),
                d.get('elapsed_ratio_multi_over_single'),
                json.dumps(mk.get('rebuilds_by_log') or {}, ensure_ascii=False),
                mk.get('rebuilds_total_log')]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v)
    autosize(ws, [14, 40, 46, 30, 30, 14])

    # ---- 4. 机审统计 ----
    by_action, trend, rules = moderation_stats()
    ws = wb.create_sheet('机审统计')
    head = sheet_header(ws, '内容机审统计（结论已落库 moderation_log）', ['机审结论', '条数'],
                        '数据直接查库（moderation_log 为只追加的事实表），可与 docs/机审效果统计-落库版.md 对照。')
    for i, (a, n) in enumerate(by_action, start=head + 1):
        ws.cell(row=i, column=1, value={'PASS': 'PASS 放行', 'REVIEW': 'REVIEW 转人工复核',
                                        'REJECT': 'REJECT 拒绝发布'}.get(a, a))
        ws.cell(row=i, column=2, value=int(n))
    if by_action:
        ch = PieChart()
        ch.title = '机审结论分布'
        ch.height, ch.width = 8, 12
        ch.add_data(Reference(ws, min_col=2, min_row=head, max_row=head + len(by_action)),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=head + len(by_action)))
        ws.add_chart(ch, f'D{head}')
    # 按天趋势
    r2 = head + len(by_action) + 3
    ws.cell(row=r2, column=1, value='按天趋势').font = HEAD_FONT
    for j, h in enumerate(['日期', 'PASS', 'REVIEW', 'REJECT'], start=1):
        c = ws.cell(row=r2 + 1, column=j, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for i, row in enumerate(trend, start=r2 + 2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v if j == 1 else int(v or 0))
    if trend:
        ch = BarChart()
        ch.type = 'col'
        ch.grouping = 'stacked'
        ch.overlap = 100
        ch.title = '机审按天趋势（堆叠）'
        ch.height, ch.width = 9, 18
        ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=r2 + 1, max_row=r2 + 1 + len(trend)),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=r2 + 2, max_row=r2 + 1 + len(trend)))
        ws.add_chart(ch, f'D{r2}')
    # 命中规则
    r3 = r2 + len(trend) + 4
    ws.cell(row=r3, column=1, value='命中的规则分布（Top 8）').font = HEAD_FONT
    for j, h in enumerate(['规则', '命中条数'], start=1):
        c = ws.cell(row=r3 + 1, column=j, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for i, (rule, n) in enumerate(rules, start=r3 + 2):
        ws.cell(row=i, column=1, value=clean(rule))
        ws.cell(row=i, column=2, value=int(n))
    autosize(ws, [26, 12, 12, 12])

    # ---- 5. 测试用例分布 ----
    dist = case_distribution()
    ws = wb.create_sheet('测试用例分布')
    head = sheet_header(ws, '自动化测试用例分布', ['类别', '条数'],
                        '数据来自 docs/测试覆盖率说明.md；合计 512 条。')
    for i, (name, n) in enumerate(dist, start=head + 1):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=n)
    if dist:
        ch = BarChart()
        ch.type = 'bar'
        ch.title = '自动化测试用例分布'
        ch.height, ch.width = 9, 18
        ch.add_data(Reference(ws, min_col=2, min_row=head, max_row=head + len(dist)),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=head + len(dist)))
        ws.add_chart(ch, f'D{head}')
    autosize(ws, [24, 12])

    # ---- 6. 缺陷统计 ----
    product, engineering, sev = defect_stats()
    ws = wb.create_sheet('缺陷统计')
    head = sheet_header(ws, '缺陷跟踪统计', ['维度', '数量'],
                        '数据来自 docs/缺陷跟踪表.md；每条都完成「发现 → 根因 → 修复 → 回归验证」四步。')
    rows = [('产品缺陷', product), ('工程与测试缺陷', engineering)]
    for i, (k, v) in enumerate(rows, start=head + 1):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    r4 = head + len(rows) + 2
    ws.cell(row=r4, column=1, value='按严重程度').font = HEAD_FONT
    for j, h in enumerate(['严重程度', '条数'], start=1):
        c = ws.cell(row=r4 + 1, column=j, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for i, (k, v) in enumerate(sorted(sev.items(), key=lambda x: -x[1]), start=r4 + 2):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    ch = BarChart()
    ch.type = 'col'
    ch.title = '缺陷分类统计'
    ch.height, ch.width = 8, 14
    ch.add_data(Reference(ws, min_col=2, min_row=head, max_row=head + len(rows)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=head + len(rows)))
    ws.add_chart(ch, f'D{head}')
    autosize(ws, [18, 12])

    # ---- 7. 索引审查 ----
    idx = load_json('数据库性能与索引审查-数据.json')
    ws = wb.create_sheet('索引审查')
    head = sheet_header(ws, '索引审查：21 种查询形态的 EXPLAIN 结果',
                        ['查询', 'table', 'type', '命中索引', '估算扫描行数', 'Extra'],
                        '在"同结构 + 真实量级"的独立库上实测（商品 5 万 / 行为 20 万 / 消息 5 万）。'
                        '数据来自 docs/数据库性能与索引审查-数据.json。')
    for i, r in enumerate(idx.get('results', []), start=head + 1):
        vals = [r.get('name'), r.get('table'), r.get('type'), r.get('key'),
                r.get('rows'), r.get('extra')]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v)
    autosize(ws, [22, 16, 12, 26, 14, 40])

    wb.save(OUT)
    return cov, pr, by_action, trend, dist, product, engineering, idx


def main():
    cov, pr, by_action, trend, dist, product, engineering, idx = build()
    print('=' * 88)
    print(f'论文图表包已生成：{OUT}')
    print(f'  覆盖率数据点      : {len(cov)}（{cov[0][2]}% → {cov[-1][2]}%）' if cov else '  覆盖率数据点: 0')
    print(f'  缓存性能对照行数  : {len(pr)}')
    print(f'  机审结论种类      : {len(by_action)}   按天趋势天数: {len(trend)}')
    print(f'  测试用例类别      : {len(dist)}   合计 {sum(n for _, n in dist)} 条')
    print(f'  缺陷统计          : 产品 {product} + 工程 {engineering}')
    print(f'  索引审查查询数    : {len(idx.get("results", []))}')
    print('=' * 88)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
